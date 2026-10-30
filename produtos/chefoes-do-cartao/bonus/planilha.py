"""Bônus 1 do ebook 'Chefões do cartão': planilha "Placar dos Chefões".
Gera produtos/chefoes-do-cartao/bonus/Placar-dos-Chefoes.xlsx (Excel e Google Planilhas).
Uso: python3 produtos/chefoes-do-cartao/bonus/planilha.py
Pesquisa e decisões: pesquisa/chefoes-planilha.md. Taxas: Banco Central, set/2026 (mesmas de contas.py)."""
import os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.label import DataLabelList

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(AQUI))
import contas  # noqa: E402  (taxas e regra oficiais do livro)

SAIDA = os.path.join(AQUI, 'Placar-dos-Chefoes.xlsx')

# ---------- identidade do canal ----------
SALVIA, CAMPO, TINTA = 'E7EFE3', 'F1F6EE', '17302A'
VERDE, VERMELHO, AMARELO = '1FBF6A', 'FF5146', 'FFCB2D'
VERDE_CLARO, VERMELHO_CLARO, AMARELO_CLARO = 'DDF7E8', 'FFE3DF', 'FFF3C4'
LETRA = 'Nunito'
BRL = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
PCT = '0.00%'

def fill(c): return PatternFill('solid', start_color=c, end_color=c)
def fonte(size=11, bold=False, cor=TINTA, italic=False): return Font(name=LETRA, size=size, bold=bold, color=cor, italic=italic)
LINHA = Side(style='thin', color='9DB59A')
BORDA = Border(left=LINHA, right=LINHA, top=LINHA, bottom=LINHA)
SOLTA = Protection(locked=False)

# ---------- tipos de dívida (ordem da lista suspensa) ----------
# (tipo, chefão do livro, taxa anual BC em fração ou None, taxa mensal fixa ou None, referência)
TIPOS = [
    ('Rotativo do cartão', 'Dragão do Rotativo', contas.ANUAL['rotativo'] / 100, None, 'BC, set/2026'),
    ('Parcelamento da fatura', 'Golem do Parcelado', contas.ANUAL['parcelado'] / 100, None, 'BC, set/2026'),
    ('Cheque especial', 'Sombra do Cheque Especial', contas.ANUAL['cheque'] / 100, None, 'BC, set/2026'),
    ('Crédito pessoal', 'Cavaleiro do Crédito Pessoal', contas.ANUAL['pessoal'] / 100, None, 'BC, set/2026'),
    ('Consignado privado (CLT)', 'Ogro do Consignado CLT', contas.ANUAL['consclt'] / 100, None, 'BC, média de ago/2026 (série 20744)'),
    ('Consignado', 'Lacaio do Consignado', contas.ANUAL['consignado'] / 100, None, 'BC, set/2026'),
    ('Financiamento de veículos', 'Tartaruga do Financiamento', contas.ANUAL['veiculos'] / 100, None, 'BC, set/2026'),
    ('Contas atrasadas (luz/água/telefone)', 'Goblin das Contas Atrasadas', None, 0.01, 'Juros de mora de 1% a.m. + multa de até 2% (ANEEL/concessionárias, via Procon-SP)'),
    ('Outra dívida', 'Chefão sem ficha', None, None, 'Sem média: digite a taxa do seu contrato em "Minha taxa"'),
]
T0 = 5                      # primeira linha de dados na aba Taxas
T1 = T0 + len(TIPOS) - 1
TIPO_RNG, CHEFAO_RNG, MENSAL_RNG = (f'Taxas!$A${T0}:$A${T1}', f'Taxas!$B${T0}:$B${T1}', f'Taxas!$D${T0}:$D${T1}')

N = 10                      # dívidas que cabem
D0 = 11                     # primeira linha de dívida em "Meus chefões"
D1 = D0 + N - 1
AVISO = ('Educativo: não é recomendação de investimento, não promete resultado e não substitui orientação profissional. '
         'Taxas médias do Banco Central de setembro de 2026 (consignado CLT: média de agosto de 2026). '
         'Confira sempre o contrato e o CET da sua dívida.')

# ---------- exemplo oficial da Ana (do livro) ----------
ANA_ORC = 600
ANA = [  # (nome, tipo, vida, mínimo, chave em contas.M)
    ('EXEMPLO · Cartão do banco', 'Parcelamento da fatura', 2000, 100, 'parcelado'),
    ('EXEMPLO · Cheque especial', 'Cheque especial', 800, 50, 'cheque'),
    ('EXEMPLO · Empréstimo do app', 'Crédito pessoal', 3000, 150, 'pessoal'),
]

def ana_meses(meses=2):
    """Avalanche da Ana, mês a mês, pela mesma regra de contas.simular(): juros, mínimos, sobra no alvo.
    Retorna lista de (pagou, vida) por mês, cada um com uma entrada por dívida."""
    ds = [dict(saldo=v, taxa=contas.M[k], minimo=m, original=v, juros=0.0) for _, _, v, m, k in ANA]
    hist = []
    for _ in range(meses):
        pago = [0.0] * len(ds)
        for d in ds:
            j = d['saldo'] * d['taxa']; j = max(0.0, min(j, d['original'] - d['juros']))  # teto não pesa no começo
            d['saldo'] += j; d['juros'] += j
        sobra = ANA_ORC
        for i, d in enumerate(ds):
            p = min(d['minimo'], d['saldo']); d['saldo'] -= p; pago[i] += p; sobra -= p
        for i in sorted(range(len(ds)), key=lambda i: -ds[i]['taxa']):
            p = min(sobra, ds[i]['saldo']); ds[i]['saldo'] -= p; pago[i] += p; sobra -= p
        hist.append(([round(p, 2) for p in pago], [round(d['saldo'], 2) for d in ds]))
    return hist

# ---------- utilidades ----------
def fundo(ws, linhas, colunas):
    """Pinta a área de uso de verde-sálvia (nada de fundo branco) e põe a letra do canal."""
    for r in range(1, linhas + 1):
        for c in range(1, colunas + 1):
            cel = ws.cell(r, c)
            cel.fill = fill(SALVIA); cel.font = fonte()
    ws.sheet_view.showGridLines = False

def celula(ws, ref, valor=None, *, size=11, bold=False, cor=TINTA, bg=None, fmt=None, h=None, v='center',
           wrap=False, borda=False, entrada=False, italic=False, merge=None):
    if merge: ws.merge_cells(merge)
    c = ws[ref]
    if valor is not None: c.value = valor
    c.font = fonte(size, bold, cor, italic)
    if entrada:
        bg = bg or CAMPO; borda = True; c.protection = SOLTA
    if bg: c.fill = fill(bg)
    if fmt: c.number_format = fmt
    c.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap)
    if borda:
        if merge:
            for linha in ws[merge]:
                for x in linha:
                    x.border = BORDA
                    if bg: x.fill = fill(bg)
        else: c.border = BORDA
    return c

def cabecalho(ws, ref, texto, merge=None, size=11, h='center'):
    c = celula(ws, ref, texto, size=size, bold=True, cor='FFFFFF', bg=TINTA, h=h, wrap=True, merge=merge)
    if merge:
        for linha in ws[merge]:
            for x in linha: x.fill = fill(TINTA)
    return c

def larguras(ws, d):
    for col, w in d.items(): ws.column_dimensions[col].width = w

def proteger(ws):
    """Proteção leve, sem senha: evita apagar fórmula sem querer. Campos de entrada ficam livres."""
    p = ws.protection
    p.sheet = True
    p.formatCells = p.formatColumns = p.formatRows = False   # False = permitido
    p.selectLockedCells = p.selectUnlockedCells = False

wb = Workbook()

# =====================================================================
# Aba: Como usar
# =====================================================================
ws = wb.active; ws.title = 'Como usar'
fundo(ws, 60, 4)
larguras(ws, {'A': 3, 'B': 6, 'C': 92, 'D': 3})
ws.sheet_properties.tabColor = VERDE
celula(ws, 'B2', 'PLACAR DOS CHEFÕES', size=26, bold=True, merge='B2:C2')
celula(ws, 'B3', 'Bônus 1 do ebook "Chefões do cartão" · Din no Chão', size=12, italic=True, merge='B3:C3')
ws.row_dimensions[2].height = 38

din = ('Oi, eu sou o Din. Esta planilha é o seu placar. Você coloca os chefões, ela mostra quem está batendo mais forte, '
       'qual atacar primeiro e se a barra de vida está caindo. Leva 10 minutos.')
celula(ws, 'B5', 'DIN', size=12, bold=True, cor='FFFFFF', bg=VERDE, h='center')
celula(ws, 'C5', din, size=12, bold=True, bg=VERDE_CLARO, wrap=True, borda=True)
ws.row_dimensions[5].height = 52

passos = [
    ('1', 'Apague o exemplo da Ana.', 'As linhas escritas "EXEMPLO" são da Ana, a personagem do livro. Selecione os campos claros com borda e aperte Delete. As fórmulas ficam.'),
    ('2', 'Monte a ficha dos seus chefões.', 'Na aba "Meus chefões": nome, tipo (lista), vida (quanto deve hoje) e o mínimo ou a parcela. A taxa ao mês aparece sozinha. Sabe a do seu contrato? Digite em "Minha taxa".'),
    ('3', 'Diga quanto sobra.', 'No topo de "Meus chefões": quanto sobra por mês para as dívidas, depois do essencial. Essencial vem antes de qualquer chefão.'),
    ('4', 'Escolha a estratégia.', 'A planilha monta a ordem de ataque da Avalanche e da Bola de Neve. O primeiro da lista é o ALVO: ele recebe toda a sobra.'),
    ('5', 'Marque o placar todo mês.', 'Na aba "Placar mensal", anote o que pagou e a vida que ficou em cada chefão. O gráfico mostra a vida total caindo.'),
    ('6', 'Siga o plano de 90 dias.', 'Uma missão por semana. Marque "Feito" quando terminar.'),
]
r = 7
celula(ws, f'B{r}', 'COMO JOGAR', bold=True, size=13, merge=f'B{r}:C{r}'); r += 1
for n, tit, txt in passos:
    celula(ws, f'B{r}', n, size=14, bold=True, cor='FFFFFF', bg=TINTA, h='center', v='top')
    c = celula(ws, f'C{r}', f'{tit} {txt}', wrap=True, v='top')
    ws.row_dimensions[r].height = 48
    r += 1

r += 1
celula(ws, f'B{r}', 'AS 3 REGRAS DO JOGO', bold=True, size=13, merge=f'B{r}:C{r}'); r += 1
for n, txt in [('1', 'Pague o essencial primeiro: comida, moradia, contas, remédio e transporte.'),
               ('2', 'Pague o mínimo de todas as dívidas, para nenhuma atrasar.'),
               ('3', 'Todo dinheiro que sobrar vai inteiro para UM chefão de cada vez: o ALVO.')]:
    celula(ws, f'B{r}', n, size=14, bold=True, cor=TINTA, bg=AMARELO, h='center')
    celula(ws, f'C{r}', txt, bold=True); ws.row_dimensions[r].height = 22; r += 1

r += 1
celula(ws, f'B{r}', 'O QUE CADA COR QUER DIZER', bold=True, size=13, merge=f'B{r}:C{r}'); r += 1
legenda = [(CAMPO, 'Campo claro com borda: é seu, preencha.', True),
           (SALVIA, 'Fundo verde-sálvia: fórmula. Não precisa mexer.', False),
           (VERMELHO, 'Vermelho: o chefão está crescendo. O mínimo não cobre o ataque dele.', False),
           (AMARELO, 'Amarelo: o ALVO, quem recebe a sobra.', False),
           (VERDE, 'Verde: chefão perdendo vida, ou fase vencida.', False)]
for cor, txt, borda in legenda:
    celula(ws, f'B{r}', None, bg=cor, borda=True)
    celula(ws, f'C{r}', txt); ws.row_dimensions[r].height = 20; r += 1

r += 1
celula(ws, f'B{r}', 'A CONTA MAIS IMPORTANTE', bold=True, size=13, merge=f'B{r}:C{r}'); r += 1
celula(ws, f'B{r}', 'Ataque do mês = vida × taxa ao mês. Se o seu pagamento é menor que o ataque, a vida do chefão sobe: você está pagando e perdendo.',
       wrap=True, bold=True, bg=VERMELHO_CLARO, borda=True, merge=f'B{r}:C{r}')
ws.row_dimensions[r].height = 34; r += 2

celula(ws, f'B{r}', 'BOM SABER', bold=True, size=13, merge=f'B{r}:C{r}'); r += 1
dicas = [
    'Taxa: digite 6,42% ou 0,0642. Se digitar 6,42 (sem o %), a planilha entende 6,42% ao mês.',
    'Onde achar os números: Registrato do Banco Central (gov.br), fatura do cartão, app do banco e Serasa.',
    'As fórmulas estão protegidas, sem senha, só para não apagar sem querer. Para mexer: Revisão > Desproteger planilha (Excel) ou Dados > Proteger páginas e intervalos (Google Planilhas).',
    'No Google Planilhas: Arquivo > Importar > Fazer upload, ou abra o arquivo pelo Google Drive e use "Abrir com Planilhas Google".',
    'A planilha não prevê a data de quitação: o mínimo do cartão muda a cada fatura. Para ver a batalha completa mês a mês, use o simulador do livro.',
]
for d in dicas:
    celula(ws, f'B{r}', '•', bold=True, h='center', v='top')
    celula(ws, f'C{r}', d, wrap=True, v='top'); ws.row_dimensions[r].height = 32; r += 1
r += 1
celula(ws, f'B{r}', AVISO, size=9, italic=True, wrap=True, merge=f'B{r}:C{r}'); ws.row_dimensions[r].height = 40
proteger(ws)

# =====================================================================
# Aba: Taxas (fonte das taxas sugeridas)
# =====================================================================
wt = wb.create_sheet('Taxas')
fundo(wt, 30, 6)
larguras(wt, {'A': 36, 'B': 30, 'C': 14, 'D': 14, 'E': 58, 'F': 3})
celula(wt, 'A1', 'TAXAS MÉDIAS DOS CHEFÕES', size=18, bold=True)
celula(wt, 'A2', 'Banco Central do Brasil, taxas médias de setembro de 2026 (divulgadas em 29/10/2026). Taxa ao mês = (1 + taxa ao ano)^(1/12) − 1, a mesma conta do livro.',
       size=10, italic=True, wrap=True, merge='A2:E2')
wt.row_dimensions[2].height = 30
for col, t in zip('ABCDE', ['Tipo de dívida', 'Chefão do livro', 'Taxa ao ano', 'Taxa ao mês', 'Fonte']):
    cabecalho(wt, f'{col}4', t)
for i, (tipo, chefao, anual, mensal, fonte_) in enumerate(TIPOS):
    r = T0 + i
    celula(wt, f'A{r}', tipo, bold=True, borda=True)
    celula(wt, f'B{r}', chefao, borda=True)
    if anual is not None:
        celula(wt, f'C{r}', anual, fmt='0.0%', borda=True)
        celula(wt, f'D{r}', f'=(1+C{r})^(1/12)-1', fmt=PCT, borda=True, bold=True)
    elif mensal is not None:
        celula(wt, f'C{r}', f'=(1+D{r})^12-1', fmt='0.0%', borda=True)
        celula(wt, f'D{r}', mensal, fmt=PCT, borda=True, bold=True)
    else:
        celula(wt, f'C{r}', None, borda=True); celula(wt, f'D{r}', None, borda=True)
    celula(wt, f'E{r}', fonte_, size=9, wrap=True, borda=True)
celula(wt, f'A{T1 + 2}', 'As taxas mudam todo mês. Atualize a coluna "Taxa ao ano" com a média nova do Banco Central (bcb.gov.br, estatísticas monetárias e de crédito), ou use a taxa do seu contrato em "Minha taxa".',
       size=10, wrap=True, merge=f'A{T1 + 2}:E{T1 + 2}')
wt.row_dimensions[T1 + 2].height = 30
celula(wt, f'A{T1 + 3}', AVISO, size=9, italic=True, wrap=True, merge=f'A{T1 + 3}:E{T1 + 3}')
wt.row_dimensions[T1 + 3].height = 40
proteger(wt)

# =====================================================================
# Aba: Meus chefões
# =====================================================================
wm = wb.create_sheet('Meus chefões', 1)
fundo(wm, 50, 15)
wm.sheet_properties.tabColor = VERMELHO
larguras(wm, {'A': 6, 'B': 30, 'C': 30, 'D': 28, 'E': 14, 'F': 11, 'G': 11, 'H': 11, 'I': 14, 'J': 14,
              'K': 14, 'L': 44, 'M': 12, 'N': 12, 'O': 3})
celula(wm, 'A1', 'MEUS CHEFÕES', size=22, bold=True, merge='A1:F1')
celula(wm, 'A2', 'Din: ninguém vence o que não enxerga. Quatro números por dívida e você sabe contra o que está lutando.',
       italic=True, merge='A2:L2')
wm.row_dimensions[1].height = 34

celula(wm, 'A4', 'Quanto sobra por mês para as dívidas (depois do essencial)?', bold=True, size=12, merge='A4:D4', h='right')
celula(wm, 'E4', ANA_ORC, fmt=BRL, bold=True, size=13, entrada=True)
celula(wm, 'F4', '← EXEMPLO da Ana. Troque pelo seu valor.', size=10, italic=True, cor='8A6A00', merge='F4:J4')
wm.row_dimensions[4].height = 26

# placar
P = [('A6:C6', 'A7:C7', 'VIDA TOTAL', f'=SUM(E{D0}:E{D1})', BRL, TINTA, 'o que você deve hoje'),
     ('D6:E6', 'D7:E7', 'ATAQUE DO MÊS', f'=SUM(J{D0}:J{D1})', BRL, VERMELHO, 'juros de 1 mês (vida × taxa)'),
     ('F6:H6', 'F7:H7', 'MÍNIMOS SOMADOS', f'=SUM(I{D0}:I{D1})', BRL, TINTA, 'para nenhuma atrasar'),
     ('I6:J6', 'I7:J7', 'SOBRA PARA O ALVO', f'=IF(E4="","",E4-F7)', BRL, VERDE, 'vai inteira para UM chefão'),
     ('K6:L6', 'K7:L7', 'DANO REAL NO MÊS', f'=IF(E4="","",E4-D7)', BRL, VERDE, 'o que pagar − o ataque do mês'),
     ('M6:N6', 'M7:N7', 'CRESCENDO', f'=COUNTIF(L{D0}:L{D1},"CRESCENDO*")', '0" chefão(ões)"', VERMELHO, 'mínimo menor que o ataque')]
for mt, mv, tit, form, fmt, cor, sub in P:
    cabecalho(wm, mt.split(':')[0], tit, merge=mt)
    c = celula(wm, mv.split(':')[0], form, size=18, bold=True, cor=cor, fmt=fmt, h='center', bg=CAMPO, borda=True, merge=mv)
    sub_ref = mv.split(':')[0].replace('7', '8'); sub_merge = mv.replace('7', '8')
    celula(wm, sub_ref, sub, size=9, italic=True, h='center', merge=sub_merge)
wm.row_dimensions[7].height = 34
celula(wm, 'I9', f'=IF(E4="","",IF(E4<F7,"Não cobre nem os mínimos: veja a Fase 4 do livro (renegociar, trocar de dívida).",""))',
       size=10, bold=True, cor=VERMELHO, merge='I9:N9', h='center')

# tabela das dívidas
heads = ['#', 'Nome da dívida (onde está)', 'Tipo (escolha na lista)', 'Chefão', 'Vida (R$)\nsaldo devedor', 'Taxa\nsugerida',
         'Minha taxa\n(opcional)', 'Taxa\nusada', 'Mínimo ou\nparcela (R$)', 'Ataque do mês\n(vida × taxa)', 'Dano do mínimo\n(mín. − ataque)',
         'Situação', 'Ordem\nAvalanche', 'Ordem\nBola de Neve']
for i, h in enumerate(heads):
    cabecalho(wm, f'{chr(65 + i)}{D0 - 1}', h)
wm.row_dimensions[D0 - 1].height = 34
valido = lambda r: f'AND(ISNUMBER(E{r}),E{r}>0,ISNUMBER(H{r}))'
VAL = f'ISNUMBER($E${D0}:$E${D1})*($E${D0}:$E${D1}>0)*ISNUMBER($H${D0}:$H${D1})'
for k in range(N):
    r = D0 + k
    celula(wm, f'A{r}', k + 1, bold=True, h='center', borda=True)
    ex = ANA[k] if k < len(ANA) else None
    celula(wm, f'B{r}', ex[0] if ex else None, entrada=True)
    celula(wm, f'C{r}', ex[1] if ex else None, entrada=True)
    celula(wm, f'D{r}', f'=IF(C{r}="","",IFERROR(INDEX({CHEFAO_RNG},MATCH(C{r},{TIPO_RNG},0)),""))', borda=True, size=10)
    celula(wm, f'E{r}', ex[2] if ex else None, fmt=BRL, entrada=True)
    celula(wm, f'F{r}', f'=IF(C{r}="","",IFERROR(IF(INDEX({MENSAL_RNG},MATCH(C{r},{TIPO_RNG},0))="","",INDEX({MENSAL_RNG},MATCH(C{r},{TIPO_RNG},0))),""))',
           fmt=PCT, borda=True)
    celula(wm, f'G{r}', None, fmt=PCT, entrada=True)
    # quem digita "6,42" sem % recebe 6,42% (ninguém paga mais de 100% ao mês)
    celula(wm, f'H{r}', f'=IF(ISNUMBER(G{r}),IF(G{r}>1,G{r}/100,G{r}),F{r})', fmt=PCT, borda=True, bold=True)
    celula(wm, f'I{r}', ex[3] if ex else None, fmt=BRL, entrada=True)
    celula(wm, f'J{r}', f'=IF({valido(r)},E{r}*H{r},"")', fmt=BRL, borda=True, bold=True, cor=VERMELHO)
    celula(wm, f'K{r}', f'=IF(AND(ISNUMBER(J{r}),ISNUMBER(I{r})),I{r}-J{r},"")', fmt=BRL, borda=True)
    celula(wm, f'L{r}', (f'=IF(NOT(ISNUMBER(E{r})),"",IF(E{r}<=0,"FASE VENCIDA",IF(NOT(ISNUMBER(H{r})),"Falta a taxa",'
                         f'IF(NOT(ISNUMBER(I{r})),"Falta o mínimo",IF(ROUND(I{r},2)<ROUND(J{r},2),"CRESCENDO: o mínimo não cobre o ataque",'
                         f'IF(ROUND(I{r},2)=ROUND(J{r},2),"Parado: o mínimo só paga os juros","Perdendo vida"))))))'),
           borda=True, bold=True, size=10)
    celula(wm, f'M{r}', f'=IF({valido(r)},1+SUMPRODUCT({VAL}*($H${D0}:$H${D1}>H{r}))+SUMPRODUCT({VAL}*($H${D0}:$H${D1}=H{r})*(ROW($H${D0}:$H${D1})<ROW(H{r}))),"")',
           fmt='0"º"', h='center', borda=True, bold=True)
    celula(wm, f'N{r}', f'=IF({valido(r)},1+SUMPRODUCT({VAL}*($E${D0}:$E${D1}<E{r}))+SUMPRODUCT({VAL}*($E${D0}:$E${D1}=E{r})*(ROW($E${D0}:$E${D1})<ROW(E{r}))),"")',
           fmt='0"º"', h='center', borda=True, bold=True)
    wm.row_dimensions[r].height = 20
TOT = D1 + 1
cabecalho(wm, f'A{TOT}', 'TOTAL', merge=f'A{TOT}:D{TOT}', h='right')
for col in 'EIJ':
    celula(wm, f'{col}{TOT}', f'=SUM({col}{D0}:{col}{D1})', fmt=BRL, bold=True, bg=TINTA, cor='FFFFFF')
for col in 'FGHKLMN':
    wm[f'{col}{TOT}'].fill = fill(TINTA)

dv_tipo = DataValidation(type='list', formula1=TIPO_RNG, allow_blank=True, showErrorMessage=True,
                         errorTitle='Tipo de dívida', error='Escolha um tipo da lista. Não achou o seu? Use "Outra dívida".')
wm.add_data_validation(dv_tipo); dv_tipo.add(f'C{D0}:C{D1}')
dv_num = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True, showErrorMessage=True,
                        errorTitle='Só números', error='Digite um valor em reais, sem letras. Ex.: 1500,00')
wm.add_data_validation(dv_num); dv_num.add(f'E{D0}:E{D1}'); dv_num.add(f'I{D0}:I{D1}'); dv_num.add('E4')
dv_taxa = DataValidation(type='decimal', operator='between', formula1='0', formula2='100', allow_blank=True, showErrorMessage=True,
                         errorTitle='Taxa ao mês', error='Digite a taxa AO MÊS. Ex.: 6,42%')
wm.add_data_validation(dv_taxa); dv_taxa.add(f'G{D0}:G{D1}')

# cores de alerta
rng_sit = f'L{D0}:L{D1}'
wm.conditional_formatting.add(rng_sit, FormulaRule(formula=[f'LEFT(L{D0},9)="CRESCENDO"'], fill=fill(VERMELHO), font=Font(name=LETRA, bold=True, color='FFFFFF')))
wm.conditional_formatting.add(rng_sit, FormulaRule(formula=[f'LEFT(L{D0},6)="Parado"'], fill=fill(AMARELO_CLARO)))
wm.conditional_formatting.add(rng_sit, FormulaRule(formula=[f'LEFT(L{D0},5)="Falta"'], fill=fill(AMARELO)))
wm.conditional_formatting.add(rng_sit, FormulaRule(formula=[f'OR(L{D0}="Perdendo vida",L{D0}="FASE VENCIDA")'], fill=fill(VERDE), font=Font(name=LETRA, bold=True, color='FFFFFF')))
wm.conditional_formatting.add(f'K{D0}:K{D1}', FormulaRule(formula=[f'AND(ISNUMBER(K{D0}),K{D0}<0)'], fill=fill(VERMELHO_CLARO), font=Font(name=LETRA, bold=True, color='D3352B')))
wm.conditional_formatting.add(f'M{D0}:N{D1}', FormulaRule(formula=[f'M{D0}=1'], fill=fill(AMARELO)))
wm.conditional_formatting.add('I7', FormulaRule(formula=['AND(ISNUMBER(I7),I7<0)'], fill=fill(VERMELHO_CLARO)))

# ordem de ataque
O0 = TOT + 2
celula(wm, f'A{O0}', 'ORDEM DE ATAQUE', size=16, bold=True, merge=f'A{O0}:F{O0}')
celula(wm, f'A{O0 + 1}', 'O primeiro da lista é o ALVO: paga o mínimo de todos e joga a sobra inteira nele. Quando ele cai, o próximo vira alvo.',
       italic=True, size=10, merge=f'A{O0 + 1}:N{O0 + 1}')
cabecalho(wm, f'A{O0 + 2}', 'AVALANCHE · maior taxa primeiro (paga menos juros)', merge=f'A{O0 + 2}:F{O0 + 2}')
cabecalho(wm, f'H{O0 + 2}', 'BOLA DE NEVE · menor vida primeiro (vitória mais cedo)', merge=f'H{O0 + 2}:N{O0 + 2}')
for col, t in zip(['A', 'B', 'C', 'E', 'F'], ['Ordem', 'Dívida', 'Chefão', 'Vida', 'Taxa']):
    celula(wm, f'{col}{O0 + 3}', t, bold=True, bg=VERDE_CLARO, h='center', borda=True)
celula(wm, f'C{O0 + 3}', None, merge=f'C{O0 + 3}:D{O0 + 3}', bg=VERDE_CLARO, borda=True)
for col, t in zip(['H', 'I', 'L', 'M', 'N'], ['Ordem', 'Dívida', 'Chefão', 'Vida', 'Taxa']):
    celula(wm, f'{col}{O0 + 3}', t, bold=True, bg=VERDE_CLARO, h='center', borda=True)
celula(wm, f'I{O0 + 3}', None, merge=f'I{O0 + 3}:K{O0 + 3}', bg=VERDE_CLARO, borda=True)

def busca(col_rank, col_val, pos):
    return f'=IF({pos}="","",IFERROR(INDEX(${col_val}${D0}:${col_val}${D1},MATCH({pos},${col_rank}${D0}:${col_rank}${D1},0)),""))'
for k in range(N):
    r = O0 + 4 + k
    bg = AMARELO if k == 0 else None
    for base, rank, cols in [('A', 'M', [('A', None), ('B', 'B'), ('C', 'D'), ('E', 'E'), ('F', 'H')]),
                             ('H', 'N', [('H', None), ('I', 'B'), ('L', 'D'), ('M', 'E'), ('N', 'H')])]:
        pos = f'${base}{r}'
        for col, src in cols:
            if src is None:
                celula(wm, f'{col}{r}', f'=IF(COUNT(${rank}${D0}:${rank}${D1})>={k + 1},{k + 1},"")', fmt='0"º"', bold=True, h='center', borda=True, bg=bg)
            else:
                fmt = BRL if src == 'E' else PCT if src == 'H' else None
                merge = f'C{r}:D{r}' if col == 'C' else f'I{r}:K{r}' if col == 'I' else None
                celula(wm, f'{col}{r}', busca(rank, src, pos), fmt=fmt, borda=True, bg=bg, bold=(k == 0), size=10, merge=merge)
    wm.row_dimensions[r].height = 19
celula(wm, f'G{O0 + 4}', '← ALVO', bold=True, size=10, cor='8A6A00')
FIM = O0 + 4 + N + 1
celula(wm, f'A{FIM}', 'Avalanche paga menos juros no total. Bola de Neve dá a primeira vitória mais cedo. Se você desanima fácil, vá de Bola de Neve; se aguenta esperar, Avalanche. As duas são muito melhores que pagar só o mínimo.',
       size=10, wrap=True, bg=VERDE_CLARO, borda=True, merge=f'A{FIM}:N{FIM}')
wm.row_dimensions[FIM].height = 32
celula(wm, f'A{FIM + 2}', AVISO, size=9, italic=True, wrap=True, merge=f'A{FIM + 2}:N{FIM + 2}')
wm.row_dimensions[FIM + 2].height = 30
wm.freeze_panes = f'A{D0}'
proteger(wm)

# =====================================================================
# Aba: Placar mensal
# =====================================================================
wp = wb.create_sheet('Placar mensal', 2)
MESES = 24
R0 = 8                       # linha do "Início"
R1 = R0 + MESES
COLS = 2 + 2 * N + 5         # A mês, B data, 2 por chefão, + 5 totais
fundo(wp, R1 + 30, COLS + 1)
wp.sheet_properties.tabColor = AMARELO
from openpyxl.utils import get_column_letter as L
cP = lambda i: L(3 + 2 * i)          # Pagou do chefão i
cV = lambda i: L(4 + 2 * i)          # Vida do chefão i
cTP, cVT, cAT, cPC, cBR = (L(3 + 2 * N + j) for j in range(5))
larguras(wp, {'A': 10, 'B': 12})
for i in range(N): larguras(wp, {cP(i): 12, cV(i): 13})
larguras(wp, {cTP: 14, cVT: 15, cAT: 15, cPC: 12, cBR: 24, L(COLS + 1): 3})

celula(wp, 'A1', 'PLACAR MENSAL', size=22, bold=True, merge='A1:H1')
celula(wp, 'A2', 'Din: todo mês, anote o que pagou em cada chefão e a vida que ficou (está no app ou na fatura). A barra de vida total tem que cair.',
       italic=True, merge=f'A2:{cV(6)}2')
wp.row_dimensions[1].height = 34
resumo = [('A3', 'A3:C3', 'A4:C4', 'VIDA NO INÍCIO', f'=IF(ISNUMBER({cVT}{R0}),{cVT}{R0},"")', BRL),
          ('D3', 'D3:F3', 'D4:F4', 'VIDA HOJE', f'=IF(COUNT({cVT}{R0}:{cVT}{R1})=0,"",INDEX({cVT}{R0}:{cVT}{R1},COUNT({cVT}{R0}:{cVT}{R1})))', BRL),
          ('G3', 'G3:I3', 'G4:I4', 'VIDA DERROTADA', f'=IF(OR(A4="",D4="",A4=0),"",1-D4/A4)', '0.0%'),
          ('J3', 'J3:L3', 'J4:L4', 'TOTAL QUE VOCÊ PAGOU', f'=SUM({cTP}{R0 + 1}:{cTP}{R1})', BRL),
          ('M3', 'M3:O3', 'M4:O4', 'ATAQUE QUE VOCÊ TOMOU', f'=SUM({cAT}{R0 + 1}:{cAT}{R1})', BRL)]
for ref, mt, mv, tit, form, fmt in resumo:
    cabecalho(wp, ref, tit, merge=mt)
    cor = VERMELHO if 'ATAQUE' in tit else VERDE if 'DERROTADA' in tit else TINTA
    celula(wp, mv.split(':')[0], form, size=16, bold=True, fmt=fmt, h='center', bg=CAMPO, borda=True, merge=mv, cor=cor)
wp.row_dimensions[4].height = 30
celula(wp, 'P4', f'=IF(G4="","",REPT("█",MAX(0,ROUND(G4*20,0)))&REPT("░",20-MAX(0,MIN(20,ROUND(G4*20,0)))))',
       size=14, cor=VERDE, merge='P4:T4')

# cabeçalhos
cabecalho(wp, f'A{R0 - 2}', 'Mês', merge=f'A{R0 - 2}:A{R0 - 1}')
cabecalho(wp, f'B{R0 - 2}', 'Data', merge=f'B{R0 - 2}:B{R0 - 1}')
for i in range(N):
    nome = f"='Meus chefões'!B{D0 + i}"
    cabecalho(wp, f'{cP(i)}{R0 - 2}', f'=IF(\'Meus chefões\'!B{D0 + i}="","Chefão {i + 1}",\'Meus chefões\'!B{D0 + i})',
              merge=f'{cP(i)}{R0 - 2}:{cV(i)}{R0 - 2}')
    celula(wp, f'{cP(i)}{R0 - 1}', 'Pagou', bold=True, bg=VERDE_CLARO, h='center', borda=True, size=10)
    celula(wp, f'{cV(i)}{R0 - 1}', 'Vida que ficou', bold=True, bg=VERDE_CLARO, h='center', borda=True, size=10)
for col, t in [(cTP, 'Total pago'), (cVT, 'VIDA TOTAL'), (cAT, 'Ataque que\nvocê tomou'), (cPC, 'Vida\nderrotada'), (cBR, 'Barra')]:
    cabecalho(wp, f'{col}{R0 - 2}', t, merge=f'{col}{R0 - 2}:{col}{R0 - 1}')
wp.row_dimensions[R0 - 2].height = 30

hist = ana_meses(2)
for m in range(MESES + 1):
    r = R0 + m
    celula(wp, f'A{r}', 'Início' if m == 0 else f'Mês {m}', bold=True, h='center', borda=True)
    celula(wp, f'B{r}', 'EXEMPLO' if m <= len(hist) else None, entrada=True, h='center', size=10,
           cor='8A6A00' if m <= len(hist) else TINTA)
    for i in range(N):
        if m == 0:
            celula(wp, f'{cP(i)}{r}', '—', h='center', borda=True, cor='8AA088')
            v = ANA[i][2] if i < len(ANA) else None
        else:
            p = hist[m - 1][0][i] if (m <= len(hist) and i < len(ANA)) else None
            celula(wp, f'{cP(i)}{r}', p, fmt=BRL, entrada=True)
            v = hist[m - 1][1][i] if (m <= len(hist) and i < len(ANA)) else None
        celula(wp, f'{cV(i)}{r}', v, fmt=BRL, entrada=True, bold=True)
    vidas = ','.join(f'{cV(i)}{r}' for i in range(N))
    pagos = ','.join(f'{cP(i)}{r}' for i in range(N))
    celula(wp, f'{cVT}{r}', f'=IF(COUNT({vidas})=0,"",SUM({vidas}))', fmt=BRL, bold=True, borda=True, size=12)
    if m == 0:
        for col in (cTP, cAT): celula(wp, f'{col}{r}', '—', h='center', borda=True, cor='8AA088')
    else:
        celula(wp, f'{cTP}{r}', f'=IF(COUNT({pagos})=0,"",SUM({pagos}))', fmt=BRL, borda=True)
        # juros reais do mês sem pedir mais nada: vida agora − vida antes + o que pagou
        celula(wp, f'{cAT}{r}', f'=IF(AND(ISNUMBER({cVT}{r}),ISNUMBER({cVT}{r - 1})),{cVT}{r}-{cVT}{r - 1}+N({cTP}{r}),"")',
               fmt=BRL, borda=True, cor=VERMELHO, bold=True)
    celula(wp, f'{cPC}{r}', f'=IF(AND(ISNUMBER({cVT}{r}),ISNUMBER({cVT}${R0}),{cVT}${R0}>0),1-{cVT}{r}/{cVT}${R0},"")', fmt='0.0%', borda=True)
    celula(wp, f'{cBR}{r}', f'=IF(ISNUMBER({cPC}{r}),REPT("█",MAX(0,MIN(20,ROUND({cPC}{r}*20,0)))),"")', borda=True, cor=VERDE, size=10)

# vida zerada = fase vencida (verde); vida que subiu em relação ao mês anterior = vermelho
for i in range(N):
    rng = f'{cV(i)}{R0 + 1}:{cV(i)}{R1}'
    a = f'{cV(i)}{R0 + 1}'; ant = f'{cV(i)}{R0}'
    wp.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({a}),{a}<=0)'], fill=fill(VERDE), font=Font(name=LETRA, bold=True, color='FFFFFF')))
    wp.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({a}),ISNUMBER({ant}),{a}>{ant})'], fill=fill(VERMELHO_CLARO), font=Font(name=LETRA, bold=True, color='D3352B')))
wp.conditional_formatting.add(f'{cVT}{R0 + 1}:{cVT}{R1}', FormulaRule(formula=[f'AND(ISNUMBER({cVT}{R0 + 1}),ISNUMBER({cVT}{R0}),{cVT}{R0 + 1}>{cVT}{R0})'], fill=fill(VERMELHO_CLARO)))

dvp = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True, showErrorMessage=True,
                     errorTitle='Só números', error='Digite um valor em reais. Ex.: 350,00')
wp.add_data_validation(dvp)
for i in range(N):
    dvp.add(f'{cP(i)}{R0 + 1}:{cP(i)}{R1}'); dvp.add(f'{cV(i)}{R0}:{cV(i)}{R1}')

nota = R1 + 2
celula(wp, f'A{nota}', ('Linha "Início": a vida de cada chefão no dia em que você começou (copie da aba Meus chefões). '
                        'As linhas EXEMPLO são a Avalanche da Ana nos 2 primeiros meses, feitas pelas contas do livro: apague e coloque as suas. '
                        '"Ataque que você tomou" = vida de agora − vida do mês anterior + o que você pagou, ou seja, os juros (e multas) que entraram no mês.'),
       size=10, wrap=True, bg=VERDE_CLARO, borda=True, merge=f'A{nota}:{cV(6)}{nota}')
wp.row_dimensions[nota].height = 46

ch = BarChart()
ch.type = 'col'; ch.style = 10
ch.title = 'Vida total dos chefões, mês a mês'
ch.y_axis.title = 'R$'; ch.y_axis.numFmt = '"R$" #,##0'
ch.y_axis.majorGridlines = None
ch.legend = None
dados = Reference(wp, min_col=3 + 2 * N + 1, min_row=R0, max_row=R1)
cats = Reference(wp, min_col=1, min_row=R0, max_row=R1)
ch.add_data(dados, titles_from_data=False); ch.set_categories(cats)
s = ch.series[0]
s.graphicalProperties = GraphicalProperties(solidFill=VERDE)
s.graphicalProperties.line.solidFill = '139256'
ch.gapWidth = 40
ch.height = 9; ch.width = 26
ch.plot_area.graphicalProperties = GraphicalProperties(solidFill=CAMPO)
ch.graphical_properties = GraphicalProperties(solidFill=SALVIA)
ch.graphical_properties.line.noFill = True
wp.add_chart(ch, f'A{nota + 2}')
celula(wp, f'A{nota + 21}', AVISO, size=9, italic=True, wrap=True, merge=f'A{nota + 21}:{cV(6)}{nota + 21}')
wp.row_dimensions[nota + 21].height = 30
wp.freeze_panes = f'C{R0}'
proteger(wp)

# =====================================================================
# Aba: Plano de 90 dias (Fase 6 do livro)
# =====================================================================
wq = wb.create_sheet('Plano de 90 dias', 3)
fundo(wq, 30, 6)
wq.sheet_properties.tabColor = TINTA
larguras(wq, {'A': 12, 'B': 22, 'C': 80, 'D': 14, 'E': 3})
celula(wq, 'A1', 'MISSÃO FINAL: 90 DIAS', size=22, bold=True, merge='A1:C1')
celula(wq, 'A2', 'Din: doze semanas, uma missão por semana. Terminou? Escolha "Feito" na última coluna.', italic=True, merge='A2:D2')
wq.row_dimensions[1].height = 34
PLANO = [
    ('Mapa do jogo', 'Puxe o Registrato, as faturas e os extratos. Preencha uma ficha por dívida.'),
    ('Placar', 'Some a vida total e o ataque do mês. Liste o essencial e descubra quanto sobra.'),
    ('Estratégia', 'Escolha Avalanche ou Bola de Neve e defina o primeiro alvo.'),
    ('Feche as portas', 'Reduza o limite do cheque especial e do cartão. Tire os cartões salvos dos apps.'),
    ('Renegociar', 'Ligue para os credores das dívidas atrasadas. Peça tudo por escrito.'),
    ('Comparar', 'Simule portabilidade e troca por dívida mais barata. Compare taxa ao mês e CET.'),
    ('Decidir', 'Feche só os acordos que cabem no bolso. Atualize as fichas.'),
    ('Gastos invisíveis', 'Corte três gastos e mande a economia para o alvo.'),
    ('Escudo', 'Comece a reserva de emergência, nem que seja com R$ 50.'),
    ('Placar de novo', 'Atualize as fichas. A vida total caiu? O ataque do mês caiu?'),
    ('Ajuste', 'Se algum chefão está crescendo, volte à Fase 4. Se tudo cai, mantenha o ritmo.'),
    ('Próximo alvo', 'Comemore cada fase vencida e escreva o plano dos próximos 90 dias.'),
]
Q0 = 7; Q1 = Q0 + len(PLANO) - 1
cabecalho(wq, 'A4', 'MISSÕES FEITAS', merge='A4:B4')
celula(wq, 'A5', f'=COUNTIF(D{Q0}:D{Q1},"Feito")&" de {len(PLANO)}"', size=16, bold=True, h='center', bg=CAMPO, borda=True, merge='A5:B5')
celula(wq, 'C5', f'=REPT("█",COUNTIF(D{Q0}:D{Q1},"Feito"))&REPT("░",{len(PLANO)}-COUNTIF(D{Q0}:D{Q1},"Feito"))', size=18, cor=VERDE)
celula(wq, 'D5', f'=IF(COUNTIF(D{Q0}:D{Q1},"Feito")={len(PLANO)},"MUNDO 1 VENCIDO!","")', bold=True, cor=VERDE, size=10, wrap=True)
wq.row_dimensions[5].height = 30
for col, t in zip('ABCD', ['Semana', 'Missão', 'O que fazer', 'Feito?']):
    cabecalho(wq, f'{col}{Q0 - 1}', t)
dvf = DataValidation(type='list', formula1='"Feito"', allow_blank=True, showErrorMessage=True,
                     errorTitle='Missão', error='Escolha "Feito" na lista, ou deixe em branco.')
wq.add_data_validation(dvf)
for k, (mis, txt) in enumerate(PLANO):
    r = Q0 + k
    celula(wq, f'A{r}', f'Semana {k + 1}', bold=True, h='center', borda=True)
    celula(wq, f'B{r}', mis, bold=True, borda=True)
    celula(wq, f'C{r}', txt, wrap=True, borda=True)
    celula(wq, f'D{r}', None, entrada=True, h='center', bold=True)
    dvf.add(f'D{r}')
    wq.row_dimensions[r].height = 30
wq.conditional_formatting.add(f'A{Q0}:D{Q1}', FormulaRule(formula=[f'$D{Q0}="Feito"'], fill=fill(VERDE_CLARO), font=Font(name=LETRA, bold=True, color='139256')))
wq.conditional_formatting.add(f'D{Q0}:D{Q1}', FormulaRule(formula=[f'$D{Q0}="Feito"'], fill=fill(VERDE), font=Font(name=LETRA, bold=True, color='FFFFFF')))
celula(wq, f'A{Q1 + 2}', 'Quando o último chefão cai, o dinheiro que ia para as dívidas continua chegando todo mês. Agora ele é seu: é o Mundo 2, O Cofre.',
       wrap=True, bg=VERDE_CLARO, borda=True, merge=f'A{Q1 + 2}:D{Q1 + 2}')
wq.row_dimensions[Q1 + 2].height = 32
celula(wq, f'A{Q1 + 4}', AVISO, size=9, italic=True, wrap=True, merge=f'A{Q1 + 4}:D{Q1 + 4}')
wq.row_dimensions[Q1 + 4].height = 30
proteger(wq)

# ---------- impressão ----------
for w in wb.worksheets:
    w.page_setup.orientation = 'landscape'
    w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 0
    w.sheet_properties.pageSetUpPr.fitToPage = True
wb.active = 0
wb.save(SAIDA)
print('ok:', SAIDA)
